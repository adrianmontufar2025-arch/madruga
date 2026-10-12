"""Lista de materiales en hoja de cálculo (.xlsx), a partir de visor/piezas.json (lo genera exportar_3d.py)."""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SALIDA = os.path.join(RAIZ, "lista_de_materiales.xlsx")

FABRICACION = {
    "P01": "Corte láser, plasma o chorro de agua desde el DXF (chapa 8 mm). Rectificar el borde de cizalla a escuadra. "
           "Templar a 52-55 HRC.",
    "P02": "Corte por hilo o láser desde el DXF (chapa 5 mm). Afilar bisel de 25-30° por un solo lado, cara plana "
           "contra el gancho. Templar a 56-58 HRC.",
    "P03": "Corte por chorro de agua o CNC (placa 8 mm). Rebaje de media madera y avellanados. Anodizado.",
    "P04": "Corte CNC o a cuchilla desde el DXF (lámina 3 mm). Avellanados para los tornillos.",
    "P05": "Mecanizado CNC desde el STEP. Anodizado.",
    "P06": "Corte láser (chapa 8 mm), escariar los agujeros. Bonificado.",
    "P07": "Corte y taladrado de pletina 16 x 6. Escariar los agujeros.",
    "P08": "Cortar el tubo a medida. Con carbono: disco diamantado y sellar el corte con epoxi.",
    "P09": "Cortar el tubo y mecanizar los extremos.",
    "P10": "Prototipo: impresión 3D en PA-CF. Serie: inyección PA6-GF30.",
    "P11": "Prototipo: impresión 3D. Serie: inyección PA6-GF30 con sobremoldeo TPE de 3 mm.",
    "P12": "Prototipo: impresión 3D. Serie: inyección PA6-GF30.",
    "P13": "Corte láser y plegado (chapa 6 mm) + ala de 7 mm.",
    "P14": "Impresión 3D, o chapa de aluminio 5052 plegada.",
    "P15": "EVA de 5 mm cortada a cuchilla + correa de velcro cosida con hebilla de liberación rápida.",
    "P16": "Abrazadera impresa en 3D + anilla de acero inoxidable comercial.",
    "P17": "Corte láser o por hilo (chapa 6 mm). Mecanizar la bisagra.",
    "P18": "Impresión 3D con recubrimiento de TPE.",
    "P19": "Patas: tubo de aluminio cortado a 311 mm. Abrazadera impresa en 3D. Puntas de acero torneadas y templadas "
           "(Ø8 x 25 mm) + disco de caucho.",
}
COMPRA = {
    "C01": "Motor sin escobillas (BLDC) de 36 V y 300-400 W, con sensores Hall. Controlador con freno eléctrico.",
    "C02": "Reductor planetario 15:1 que acople al eje del motor.",
    "C03": "Husillo de bolas SFU1605 (Ø16 mm, paso 5 mm) con soporte de rodamiento. Largo según el plano de conjunto.",
    "C04": "Tuerca SFU1605 + horquilla de tiro mecanizada.",
    "C05": "Batería de herramienta de 36 V. Con 2,5 Ah pesa unos 0,9 kg y con 4 Ah unos 1,3 kg. Va al cinturón con "
           "cable espiral.",
    "C06": "Perno con hombro ISO 7379 de M12 + arandelas Belleville (mantienen el juego de 0,05-0,1 mm).",
    "C07": "Pasador Ø6 + separador + resorte de torsión de 0,3 N·m.",
    "C08": "Arnés simple de desbrozadora, con mosquetón de liberación rápida.",
    "C09": "Palanca de pulgar + cable de freno de bicicleta con funda.",
}

FUENTE = "Arial"
GRIS = PatternFill("solid", start_color="E7E9EC")
FINO = Side(style="thin", color="B5BAC1")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def main():
    with open(os.path.join(RAIZ, "visor", "piezas.json"), encoding="utf-8") as f:
        datos = json.load(f)
    piezas = {p["codigo"]: p for p in datos["piezas"]}
    piezas.setdefault("C09", dict(codigo="C09", nombre="Botón de pulgar + cable de la pata P19", grupo="Brazo",
                                  material="Comercial", cantidad=1, masa_g=0, plano=False))
    codigos = sorted((c for c in piezas if c != "C05"), key=lambda c: (c[0] != "P", c))
    planos = set(os.listdir(os.path.join(RAIZ, "planos_dxf")))
    steps = set(os.listdir(os.path.join(RAIZ, "modelo_3d", "piezas_step")))

    def archivo(lista, cod):
        return next((n for n in sorted(lista) if n.startswith(cod + "_")), "—")

    wb = Workbook()
    ws = wb.active
    ws.title = "Lista de materiales"
    ws["A1"] = "Tijera de caña con apoyo de antebrazo · Lista de materiales (Rev. B, conceptual)"
    ws["A1"].font = Font(name=FUENTE, size=14, bold=True)
    ws["A2"] = ("Masas calculadas del modelo 3D (fuente/modelo.py). Diseño para prototipo: confirmar la fuerza de corte "
                "con un ensayo antes de fabricar las piezas de acero.")
    ws["A2"].font = Font(name=FUENTE, size=10, italic=True, color="555B63")

    cab = ["Código", "Denominación", "Grupo", "Material / especificación", "Cant.", "Masa unit. (g)", "Masa total (g)",
           "Cómo se fabrica / qué comprar", "Plano DXF", "Modelo 3D (STEP)"]
    fila_cab = 4
    for j, t in enumerate(cab, 1):
        c = ws.cell(row=fila_cab, column=j, value=t)
        c.font = Font(name=FUENTE, size=10, bold=True)
        c.fill = GRIS
        c.border = BORDE
        c.alignment = Alignment(vertical="center", wrap_text=True)

    def fila(r, p, nota):
        cod = p["codigo"]
        es_pieza = cod.startswith("P")
        valores = [cod, p["nombre"], p["grupo"], p["material"], p["cantidad"], p["masa_g"], f"=E{r}*F{r}", nota,
                   archivo(planos, cod) if es_pieza else "comercial", archivo(steps, cod) if es_pieza else "—"]
        for j, v in enumerate(valores, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name=FUENTE, size=10)
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=j in (2, 4, 8))
            if j in (6, 7):
                c.number_format = "#,##0"

    r = fila_cab + 1
    for cod in codigos:
        p = piezas[cod]
        nota = FABRICACION.get(cod) or COMPRA.get(cod, "")
        if cod == "C09":
            nota += " Masa incluida en la fila de tornillería."
        fila(r, p, nota)
        r += 1
    fila(r, dict(codigo="—", nombre="Tornillería, cables y electrónica (controlador, sensor Hall del gatillo, "
                 "sensor de presencia)", grupo="Varios", material="Comercial", cantidad=1,
                 masa_g=datos["tornilleria_cables_g"]), "Estimado; no está dibujado en el modelo 3D.")
    ws.cell(row=r, column=9, value="—")
    ultima = r
    r += 1
    ws.cell(row=r, column=6, value="Masa sin batería (g)").font = Font(name=FUENTE, size=10, bold=True)
    c = ws.cell(row=r, column=7, value=f"=SUM(G{fila_cab + 1}:G{ultima})")
    c.font = Font(name=FUENTE, size=10, bold=True)
    c.number_format = "#,##0"
    c.border = BORDE
    fila_sin = r
    r += 2
    ws.cell(row=r, column=1, value="Al cinturón (no pesa en el brazo)").font = Font(name=FUENTE, size=10, bold=True)
    r += 1
    bat = dict(piezas["C05"])
    bat["masa_g"] = 900
    fila(r, bat, COMPRA["C05"])
    ws.cell(row=r, column=6).comment = None
    fila_bat = r
    r += 1
    ws.cell(row=r, column=6, value="Total con batería (g)").font = Font(name=FUENTE, size=10, bold=True)
    c = ws.cell(row=r, column=7, value=f"=G{fila_sin}+G{fila_bat}")
    c.font = Font(name=FUENTE, size=10, bold=True)
    c.number_format = "#,##0"
    c.border = BORDE
    r += 2
    notas = [
        "Notas",
        "Masa de la batería: 900 g para 2,5 Ah (estimado de catálogo). Con 4 Ah, unos 1.300 g.",
        "Los DXF están a escala 1:1: las piezas de chapa (P01, P02, P03, P04, P06, P07, P13, P17) van directo a corte;"
        " antes, borrar las capas COTAS, TEXTO, EJES y MARCO.",
        "Las piezas plásticas (P10, P11, P12, P14, P16, P18) se pueden imprimir en 3D desde el STEP para el prototipo.",
    ]
    for k, t in enumerate(notas):
        c = ws.cell(row=r + k, column=1, value=t)
        c.font = Font(name=FUENTE, size=10, bold=(k == 0))

    for col, ancho in zip("ABCDEFGHIJ", (8, 34, 16, 40, 6, 12, 12, 58, 42, 42)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = ws.cell(row=fila_cab + 1, column=3)
    ws.print_title_rows = f"{fila_cab}:{fila_cab}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(SALIDA)
    print(SALIDA)


if __name__ == "__main__":
    main()
